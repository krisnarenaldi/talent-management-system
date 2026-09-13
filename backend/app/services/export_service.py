"""Business logic untuk export data."""
from datetime import datetime, timedelta
from io import BytesIO
from typing import Optional

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from sqlalchemy.orm import Session, joinedload

from app.models.application import Application, STAGE_NAMES
from app.models.candidate import Candidate, CandidateDocument
from app.models.position import Position


class ExportError(Exception):
    pass


def _get_excel_response_for_candidates(
    db: Session,
    position_id: Optional[str] = None,
    completeness_status: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
) -> bytes:
    """
    Generate Excel file with candidates data.
    Returns bytes of the Excel file.
    """
    # Build query with joins
    query = (
        db.query(Candidate)
        .options(
            joinedload(Candidate.educations),
            joinedload(Candidate.experiences),
            joinedload(Candidate.documents),
        )
        .filter(Candidate.is_deleted == False)
    )

    # Apply filters
    if position_id:
        # Candidates yang punya application ke posisi tertentu
        query = query.join(Application, Candidate.id == Application.candidate_id).filter(
            Application.position_id == position_id
        )

    if completeness_status:
        query = query.filter(Candidate.completeness_status == completeness_status)

    if start_date:
        query = query.filter(Candidate.created_at >= start_date)
    if end_date:
        query = query.filter(Candidate.created_at <= end_date)

    candidates = query.all()

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"

    # Header row
    headers = [
        "No",
        "ID",
        "Nama",
        "Email",
        "Phone",
        "NIK (KTP)",
        "TTL",
        "Jenis Kelamin",
        "Golongan Darah",
        "Domisili",
        "Sumber Channel",
        "Gaji Saat Ini",
        "Gaji Diharapkan",
        "Notice Period (hari)",
        "Status Kelengkapan",
        "Status Kontak",
        "Catatan",
        "Tanggal Dibuat",
        "Tanggal Diperbarui",
        "Dokumen",
        "Pendidikan Terakhir",
        "Pengalaman Kerja Terakhir",
    ]

    ws.append(headers)

    # Style header
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for idx, candidate in enumerate(candidates, 1):
        # Get last education from joinedload relationship
        last_education = None
        if candidate.educations:
            last_education = sorted(
                candidate.educations,
                key=lambda x: x.graduation_year or 0,
                reverse=True,
            )[0]

        # Get last experience from joinedload relationship
        last_experience = None
        if candidate.experiences:
            last_experience = sorted(
                candidate.experiences,
                key=lambda x: x.start_date or datetime.min.date(),
                reverse=True,
            )[0]

        # Count documents
        doc_count = len(candidate.documents)

        row_data = [
            idx,
            str(candidate.id),
            candidate.full_name or "",
            candidate.email or "",
            candidate.phone or "",
            candidate.identity_no or "",
            f"{candidate.birth_place or ''}, {candidate.birth_date or ''}",
            candidate.gender or "",
            candidate.blood_type or "",
            candidate.domicile or "",
            candidate.source_channel or "",
            str(candidate.current_salary or ""),
            str(candidate.expected_salary or ""),
            str(candidate.notice_period_days or ""),
            candidate.completeness_status,
            candidate.contact_status,
            candidate.notes or "",
            candidate.created_at.strftime("%Y-%m-%d %H:%M:%S") if candidate.created_at else "",
            candidate.updated_at.strftime("%Y-%m-%d %H:%M:%S") if candidate.updated_at else "",
            f"{doc_count} dokumen",
            f"{last_education.institution if last_education else ''} ({last_education.graduation_year if last_education else ''})",
            f"{last_experience.company_name if last_experience else ''} - {last_experience.job_title if last_experience else ''}",
        ]

        ws.append(row_data)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to bytes buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.read()


def _get_excel_response_for_pipeline(
    db: Session,
    status_filter: Optional[str] = None,
    current_stage: Optional[str] = None,
    position_id: Optional[str] = None,
    start_date: Optional[datetime] = None,
    end_date: Optional[datetime] = None,
) -> bytes:
    """
    Generate Excel file with pipeline data.
    Returns bytes of the Excel file.
    """
    # Build query with joins
    query = (
        db.query(Application)
        .options(
            joinedload(Application.candidate),
            joinedload(Application.position).joinedload(Position.client),
            joinedload(Application.recruiter),
            joinedload(Application.stage_histories),
        )
    )

    # Apply filters
    if status_filter:
        query = query.filter(Application.status == status_filter)
    if current_stage:
        query = query.filter(Application.current_stage == current_stage)
    if position_id:
        query = query.filter(Application.position_id == position_id)

    if start_date:
        query = query.filter(Application.created_at >= start_date)
    if end_date:
        query = query.filter(Application.created_at <= end_date)

    applications = query.all()

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Pipeline"

    # Header row
    headers = [
        "No",
        "ID Lamaran",
        "Nama Kandidat",
        "Email Kandidat",
        "Phone Kandidat",
        "Posisi",
        "Client",
        "Stage Saat Ini",
        "Status Lamaran",
        "Tanggal Dibuat",
        "Tanggal Diperbarui",
        "Terakhir Update Stage",
        "Rekruter",
        "Hasil Terakhir",
    ]

    ws.append(headers)

    # Style header
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for idx, app in enumerate(applications, 1):
        # Get last stage history
        last_stage_history = None
        if app.stage_histories:
            last_stage_history = sorted(
                app.stage_histories,
                key=lambda x: x.created_at or datetime.min.replace(tzinfo=app.created_at.tzinfo),
            )[-1]

        row_data = [
            idx,
            str(app.id),
            app.candidate_name or "",
            app.candidate.email if app.candidate else "",
            app.candidate.phone if app.candidate else "",
            app.position_title if app.position else "",
            app.client_name if app.position and app.position.client else "",
            app.current_stage,
            app.status,
            app.created_at.strftime("%Y-%m-%d %H:%M:%S") if app.created_at else "",
            app.updated_at.strftime("%Y-%m-%d %H:%M:%S") if app.updated_at else "",
            last_stage_history.created_at.strftime("%Y-%m-%d %H:%M:%S") if last_stage_history and last_stage_history.created_at else "",
            app.recruiter_name if app.recruiter else "",
            last_stage_history.result if last_stage_history else "",
        ]

        ws.append(row_data)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to bytes buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.read()


def _get_excel_response_for_incomplete_candidates(
    db: Session,
    position_id: Optional[str] = None,
) -> bytes:
    """
    Generate Excel file with candidates that have incomplete data.
    Returns bytes of the Excel file.
    """
    # Build query with joins
    query = (
        db.query(Candidate)
        .options(
            joinedload(Candidate.documents),
            joinedload(Candidate.educations),
            joinedload(Candidate.experiences),
        )
        .filter(Candidate.is_deleted == False)
        .filter(Candidate.completeness_status == "belum_lengkap")
    )

    # Apply position filter
    if position_id:
        query = query.join(Application, Candidate.id == Application.candidate_id).filter(
            Application.position_id == position_id
        )

    candidates = query.all()

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Incomplete Candidates"

    # Header row
    headers = [
        "No",
        "ID",
        "Nama",
        "Email",
        "Phone",
        "NIK (KTP)",
        "Dokumen Lengkap",
        "Dokumen Upload",
        "Dokumen KTP",
        "Dokumen Ijazah",
        "Dokumen Transkrip",
        "Dokumen CV_asli",
        "Dokumen Sertifikat",
        "Dokumen BPJS_TK",
        "Dokumen BPJS_Kesehatan",
        "Dokumen NPWP",
        "Catatan",
        "Tanggal Dibuat",
        "Tanggal Diperbarui",
    ]

    ws.append(headers)

    # Style header
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # Data rows
    for idx, candidate in enumerate(candidates, 1):
        # Check which required documents are present
        docs = {doc.doc_type: doc for doc in candidate.documents}

        required_docs = {
            "Foto": False,
            "KTP": False,
            "Ijazah": False,
            "Transkrip": False,
            "CV_asli": False,
            "Sertifikat": False,
            "BPJS_TK": False,
            "BPJS_Kesehatan": False,
            "NPWP": False,
        }

        for doc_type in required_docs:
            if doc_type in docs:
                required_docs[doc_type] = True

        row_data = [
            idx,
            str(candidate.id),
            candidate.full_name or "",
            candidate.email or "",
            candidate.phone or "",
            candidate.identity_no or "",
            f"{sum(required_docs.values())}/{len(required_docs)} dokumen lengkap",
            "✓" if required_docs.get("Foto", False) else "✗",
            "✓" if required_docs.get("KTP", False) else "✗",
            "✓" if required_docs.get("Ijazah", False) else "✗",
            "✓" if required_docs.get("Transkrip", False) else "✗",
            "✓" if required_docs.get("CV_asli", False) else "✗",
            "✓" if required_docs.get("Sertifikat", False) else "✗",
            "✓" if required_docs.get("BPJS_TK", False) else "✗",
            "✓" if required_docs.get("BPJS_Kesehatan", False) else "✗",
            "✓" if required_docs.get("NPWP", False) else "✗",
            candidate.notes or "",
            candidate.created_at.strftime("%Y-%m-%d %H:%M:%S") if candidate.created_at else "",
            candidate.updated_at.strftime("%Y-%m-%d %H:%M:%S") if candidate.updated_at else "",
        ]

        ws.append(row_data)

    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Save to bytes buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer.read()